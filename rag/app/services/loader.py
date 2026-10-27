"""多格式文档解析。

职责边界：只把「一个文件」变成「若干 Section」，不负责切片（chunker.py 负责）。
这样新增格式时只改这里，切片策略调整时也不用动解析代码。

支持格式与实现方式：
- .md/.markdown  直接读文本，按 ATX 标题（# / ##）切分章节
- .txt           整篇作为一个 Section
- .pdf           pypdf 逐页 extract_text，页码存进 Section.page
- .docx          python-docx，按段落样式识别标题，表格按行拼成文本
- .doc           老二进制格式，Python 无纯库可读，依赖 antiword 或 soffice 转换
- .xmind         xmind 本质是 zip，读 content.json（新版）或 content.xml（旧版），
                 按主题树把叶子节点拼成「父 / 子 / 孙」的文本行
"""

import json
import re
import shutil
import subprocess
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree

from ..errors import DocumentParseError, UnsupportedFormatError

# 允许上传的扩展名白名单；api 层用它做前置校验，pipeline 用它做二次校验
SUPPORTED_SUFFIXES = {".md", ".markdown", ".txt", ".pdf", ".doc", ".docx", ".xmind"}

# 匹配 Markdown 的 1~6 级标题，group(2) 是标题文本
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")


@dataclass
class Section:
    """解析后的章节：一个标题下的正文。

    heading 为空表示该段没有标题（纯 txt、doc 转换结果常见）。
    page 只对 PDF 有意义，用来在引用里展示页码。
    """

    heading: str
    text: str
    page: int | None = None


def load(path: Path) -> tuple[str, list[Section]]:
    """按扩展名分发解析，返回 (格式名, 章节列表)。

    每个 _load_xxx 内部保证返回非空列表，解析不出文本时抛 DocumentParseError，
    避免后续生成出零个切片、库里留下"幽灵文档"。
    """
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown"}:
        return "md", _load_markdown(path)
    if suffix == ".txt":
        return "txt", _load_text(path)
    if suffix == ".pdf":
        return "pdf", _load_pdf(path)
    if suffix == ".docx":
        return "docx", _load_docx(path)
    if suffix == ".doc":
        return "doc", _load_doc(path)
    if suffix == ".xmind":
        return "xmind", _load_xmind(path)
    raise UnsupportedFormatError(f"不支持的文件类型：{suffix or '无扩展名'}")


def doc_tool_available() -> bool:
    """老版 .doc 依赖外部命令，/health 用它提前暴露"环境不支持"的情况。"""
    return bool(shutil.which("antiword") or shutil.which("soffice"))


# ---------------------------------------------------------------- 纯文本类


def _load_text(path: Path) -> list[Section]:
    """txt 不切章节，整篇作为一个 Section 交给 chunker 按长度切。"""
    text = _read_text(path).strip()
    if not text:
        raise DocumentParseError("文档内容为空")
    return [Section(heading="", text=text)]


def _load_markdown(path: Path) -> list[Section]:
    """按标题切分章节，标题本身作为元数据，不进正文。

    这样切出的每个 chunk 都能带上 heading，回答时可以说清楚
    "这段话出自哪一节"，比整篇切块更容易溯源。
    """
    sections: list[Section] = []
    heading = ""
    buffer: list[str] = []

    def flush() -> None:
        """把当前缓冲区收成一个 Section；空正文（只有标题没有内容）则跳过。"""
        body = "\n".join(buffer).strip()
        if body:
            sections.append(Section(heading=heading, text=body))
        buffer.clear()

    for line in _read_text(path).splitlines():
        match = _HEADING_RE.match(line)
        if match:
            # 遇到新标题：先结算上一节，再更新当前标题
            flush()
            heading = match.group(2)
        else:
            buffer.append(line)
    flush()  # 收尾：最后一节没有后续标题触发 flush
    return _ensure(sections)


# ---------------------------------------------------------------- PDF


def _load_pdf(path: Path) -> list[Section]:
    """逐页提取文本。pypdf 纯 Python、无需系统依赖，故延迟导入以加快冷启动。"""
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise DocumentParseError("缺少依赖 pypdf，请先 uv add pypdf") from exc

    sections: list[Section] = []
    # start=1 让页码从 1 开始，和 PDF 阅读器显示一致
    for page_no, page in enumerate(PdfReader(str(path)).pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            sections.append(Section(heading=f"第 {page_no} 页", text=text, page=page_no))

    # 整篇都提不出文字，通常是扫描件，需要 OCR——这里明确报错而不是静默入库
    return _ensure(sections, "PDF 未提取到文本，可能是扫描件，需要先做 OCR")


# ---------------------------------------------------------------- Word


def _load_docx(path: Path) -> list[Section]:
    """按段落样式切章节，表格追加到当前章节末尾。

    Word 里标题的样式名通常是 "Heading 1"/"标题 1"，中英文版本不同，
    所以用 startswith(("heading", "标题")) 兼容两种情况。
    """
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise DocumentParseError("缺少依赖 python-docx，请先 uv add python-docx") from exc

    document = Document(str(path))
    sections: list[Section] = []
    heading = ""
    buffer: list[str] = []

    def flush() -> None:
        body = "\n".join(buffer).strip()
        if body:
            sections.append(Section(heading=heading, text=body))
        buffer.clear()

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue  # 空段落（Word 里很常见）直接跳过
        style = (para.style.name or "") if para.style is not None else ""
        if style.lower().startswith(("heading", "标题")):
            flush()
            heading = text
        else:
            buffer.append(text)

    # 表格在 document.paragraphs 里看不到，单独遍历；
    # 每行用 " | " 连接，让模型能看懂行列关系
    for table in document.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        rows = [row for row in rows if row.strip(" |")]  # 丢掉整行空白
        if rows:
            buffer.append("\n".join(rows))

    flush()
    return _ensure(sections, "docx 中未解析到文本")


def _load_doc(path: Path) -> list[Section]:
    """老版 .doc（OLE 二进制）没有可用的纯 Python 解析器，只能借助外部命令。

    优先级：
    1. antiword：专读 .doc，输出纯文本，最轻量；
    2. soffice（LibreOffice）：先转成 .docx 再复用 _load_docx，兼容性更好但更慢；
    3. 都没有则明确报错，提示用户另存为 .docx——绝不静默返回空文本。
    """
    if shutil.which("antiword"):
        result = subprocess.run(["antiword", str(path)], capture_output=True, check=False)
        # antiword 输出默认不是 UTF-8，用 replace 兜底，避免编码异常中断整个请求
        text = result.stdout.decode("utf-8", "replace").strip()
        if result.returncode == 0 and text:
            return [Section(heading="", text=text)]

    if shutil.which("soffice"):
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "docx", "--outdir", str(path.parent), str(path)],
            capture_output=True,
            timeout=120,  # 转换可能卡住，给个上限
            check=False,
        )
        converted = path.parent / f"{path.stem}.docx"
        if converted.exists():
            try:
                return _load_docx(converted)
            finally:
                # 转换产物是临时文件，解析完立即删除，避免污染 uploads 目录
                converted.unlink(missing_ok=True)

    raise DocumentParseError(
        ".doc 是二进制格式，需要 antiword 或 LibreOffice(soffice) 才能解析，建议先另存为 .docx"
    )


# ---------------------------------------------------------------- XMind


def _load_xmind(path: Path) -> list[Section]:
    """xmind 文件是 zip 包：新版 XMind 2020+ 用 content.json，旧版用 content.xml。"""
    try:
        archive = zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        raise DocumentParseError("xmind 文件损坏，或不是合法的 zip 包") from exc

    with archive:
        names = set(archive.namelist())
        if "content.json" in names:
            data = json.loads(archive.read("content.json").decode("utf-8", "replace"))
            return _ensure(_xmind_from_json(data), "xmind 中未解析到主题")
        if "content.xml" in names:
            return _ensure(_xmind_from_xml(archive.read("content.xml")), "xmind 中未解析到主题")

    raise DocumentParseError("xmind 中找不到 content.json 或 content.xml")


def _xmind_from_json(data: object) -> list[Section]:
    """解析新版 content.json。

    结构大致为：[{ "title": "画布1", "rootTopic": {...} }, ...]
    一个 xmind 文件可以有多张画布，每张画布对应一个 Section。
    """
    sheets = data if isinstance(data, list) else [data]
    sections: list[Section] = []
    for sheet in sheets:
        if not isinstance(sheet, dict):
            continue  # 防御异常结构，跳过而不是整个文件失败
        lines: list[str] = []
        _walk_json(sheet.get("rootTopic") or {}, [], lines)
        if lines:
            sections.append(Section(heading=(sheet.get("title") or "Sheet").strip(), text="\n".join(lines)))
    return sections


def _walk_json(topic: dict, path: list[str], out: list[str]) -> None:
    """深度优先遍历主题树，只在叶子节点输出整条路径。

    例：根 -> 秘密道具 -> 竹蜻蜓，输出 "根 / 秘密道具 / 竹蜻蜓"。
    中间节点不单独输出，否则同一条路径会被重复写入很多次。
    """
    title = (topic.get("title") or "").strip()
    current = path + [title] if title else path
    # xmind 的附属子主题挂在 children.attached 下
    children = (topic.get("children") or {}).get("attached") or []
    if not children:
        if current:
            out.append(" / ".join(current))
        return
    for child in children:
        _walk_json(child, current, out)


def _xmind_from_xml(raw: bytes) -> list[Section]:
    """解析旧版 content.xml（带命名空间的 XML）。"""
    root = ElementTree.fromstring(raw)
    sections: list[Section] = []
    for sheet in (el for el in root.iter() if _local(el.tag) == "sheet"):
        title = ""
        root_topic = None
        for child in sheet:
            if _local(child.tag) == "title":
                title = (child.text or "").strip()
            elif _local(child.tag) == "topic":
                root_topic = child
        if root_topic is None:
            continue
        lines: list[str] = []
        _walk_xml(root_topic, [], lines)
        if lines:
            sections.append(Section(heading=title or "Sheet", text="\n".join(lines)))
    return sections


def _walk_xml(topic: ElementTree.Element, path: list[str], out: list[str]) -> None:
    """XML 版的主题树遍历，逻辑与 _walk_json 一致。

    XML 层级比 JSON 多两层包装：<children><topics><topic>，
    所以要先穿过 children/topics 才能拿到子主题。
    """
    name = ""
    children: list[ElementTree.Element] = []
    for child in topic:
        tag = _local(child.tag)
        if tag == "title" and not name:
            name = (child.text or "").strip()
        elif tag == "children":
            for block in child:
                if _local(block.tag) == "topics":
                    children.extend(el for el in block if _local(el.tag) == "topic")

    current = path + [name] if name else path
    if not children:
        if current:
            out.append(" / ".join(current))
        return
    for child in children:
        _walk_xml(child, current, out)


# ---------------------------------------------------------------- 工具函数


def _local(tag: str) -> str:
    """去掉 XML 命名空间，例如 "{http://xmind.net/...}topic" -> "topic"。

    不同版本 xmind 的命名空间 URI 不一样，按本地名匹配最稳。
    """
    return tag.rsplit("}", 1)[-1]


def _read_text(path: Path) -> str:
    """统一按 UTF-8 读文本；errors="replace" 保证个别乱码不会让整个请求失败。"""
    return path.read_text(encoding="utf-8", errors="replace")


def _ensure(sections: list[Section], message: str = "文档内容为空") -> list[Section]:
    """统一空结果校验：解析不出内容就报 400，避免写入零个切片。"""
    if not sections:
        raise DocumentParseError(message)
    return sections
