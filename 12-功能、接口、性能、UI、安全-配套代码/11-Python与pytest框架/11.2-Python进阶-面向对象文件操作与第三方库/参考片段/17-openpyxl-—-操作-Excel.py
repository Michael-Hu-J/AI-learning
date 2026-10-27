from openpyxl import load_workbook

# 读 Excel
wb = load_workbook("测试用例.xlsx")
sheet = wb.active
for row in sheet.iter_rows(min_row=2, values_only=True):
    print(f"用例：{row[0]}，预期：{row[1]}")

# 写结果
sheet["D1"] = "结果"
for i, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
    sheet[f"D{i}"] = "通过"
wb.save("测试用例_结果.xlsx")
